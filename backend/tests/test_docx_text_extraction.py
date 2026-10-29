import hashlib
import io
import os
import uuid
import zipfile
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock

import pytest
import docx
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from app.config import Settings
from app.models.extraction import BlockType, ExtractionMethod, ResumeExtraction
from app.models.resume import (
    ArtifactFormat,
    ArtifactValidationStatus,
    ResumeArtifact,
)
from app.schemas.extraction import ExtractionResult
from app.services.docx_text_extraction import (
    ArtifactIntegrityError,
    DOCXExtractionFailedError,
    DOCXExtractionOutcome,
    DOCXTextLimitExceededError,
    DOCXTextUnavailableError,
    extract_docx_text,
)
from app.services.extraction_persistence import persist_extraction


def _add_hyperlink(paragraph, url: str, text: str):
    """Helper to safely insert a w:hyperlink element into a paragraph with visible text."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = parse_xml(f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{r_id}"/>')
    new_run = parse_xml(f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:rPr/><w:t>{text}</w:t></w:r>')
    hyperlink.append(new_run)
    paragraph._element.append(hyperlink)


def _setup_quarantined_docx(
    tmp_path: Path,
    docx_bytes: bytes,
    validation_status: ArtifactValidationStatus = ArtifactValidationStatus.VALIDATED,
    detected_format: ArtifactFormat = ArtifactFormat.docx,
    quarantine_key_override: str = None,
    size_override: int = None,
    hash_override: str = None,
) -> tuple[ResumeArtifact, Settings]:
    """Helper to setup quarantine directory, write DOCX file, and build ResumeArtifact."""
    quarantine_dir = tmp_path / "quarantine"
    docx_dir = quarantine_dir / "docx"
    docx_dir.mkdir(parents=True, exist_ok=True)

    art_id = uuid.uuid4()
    docx_filename = f"{art_id}.docx"
    file_path = docx_dir / docx_filename
    file_path.write_bytes(docx_bytes)

    actual_size = len(docx_bytes)
    actual_hash = hashlib.sha256(docx_bytes).hexdigest().lower()

    quarantine_key = quarantine_key_override if quarantine_key_override is not None else f"docx/{docx_filename}"
    size_bytes = size_override if size_override is not None else actual_size
    sha256_hash = hash_override if hash_override is not None else actual_hash

    artifact = ResumeArtifact(
        artifact_id=art_id,
        original_filename="candidate_resume.docx",
        declared_mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        detected_format=detected_format,
        size_bytes=size_bytes,
        sha256_hash=sha256_hash,
        validation_status=validation_status,
        quarantine_key=quarantine_key,
    )

    settings = Settings(
        database_url="postgresql+psycopg://dummy:dummy@localhost:5432/dummy",
        resume_quarantine_directory=str(quarantine_dir),
        max_docx_extracted_chars=500000,
    )

    return artifact, settings


# =========================================================================
# 1. Interleaved paragraph-table-paragraph order is preserved
# =========================================================================

def test_interleaved_body_order_preserved(tmp_path):
    doc = Document()
    doc.add_paragraph("Paragraph 0 - Before Table")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cell 0,0"
    table.cell(0, 1).text = "Cell 0,1"
    table.cell(1, 0).text = "Cell 1,0"
    table.cell(1, 1).text = "Cell 1,1"

    doc.add_paragraph("Paragraph 1 - After Table")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(tmp_path, docx_bytes)
    outcome = extract_docx_text(artifact, settings)

    assert isinstance(outcome, DOCXExtractionOutcome)
    assert outcome.review_required is False
    assert len(outcome.blocks) == 6

    # Verify contiguous sequence indexes
    assert [b.sequence_index for b in outcome.blocks] == [0, 1, 2, 3, 4, 5]

    # Check exact interleaved order and structural locations
    assert outcome.blocks[0].raw_text == "Paragraph 0 - Before Table"
    assert outcome.blocks[0].structural_location == "body/paragraph/0"
    assert outcome.blocks[0].block_type == BlockType.PARAGRAPH

    assert outcome.blocks[1].raw_text == "Cell 0,0"
    assert outcome.blocks[1].structural_location == "body/table/0/row/0/cell/0"
    assert outcome.blocks[1].block_type == BlockType.TABLE_CELL

    assert outcome.blocks[2].raw_text == "Cell 0,1"
    assert outcome.blocks[2].structural_location == "body/table/0/row/0/cell/1"
    assert outcome.blocks[2].block_type == BlockType.TABLE_CELL

    assert outcome.blocks[3].raw_text == "Cell 1,0"
    assert outcome.blocks[3].structural_location == "body/table/0/row/1/cell/0"
    assert outcome.blocks[3].block_type == BlockType.TABLE_CELL

    assert outcome.blocks[4].raw_text == "Cell 1,1"
    assert outcome.blocks[4].structural_location == "body/table/0/row/1/cell/1"
    assert outcome.blocks[4].block_type == BlockType.TABLE_CELL

    assert outcome.blocks[5].raw_text == "Paragraph 1 - After Table"
    assert outcome.blocks[5].structural_location == "body/paragraph/1"
    assert outcome.blocks[5].block_type == BlockType.PARAGRAPH


# =========================================================================
# 2. Paragraphs, headings, and list items receive appropriate block types
# =========================================================================

def test_paragraph_heading_list_classification(tmp_path):
    doc = Document()
    doc.add_heading("Senior Architect Resume", level=1)
    doc.add_paragraph("Summary of qualifications and experience.")
    doc.add_paragraph("Built cloud platform with 99.99% reliability", style="List Bullet")
    doc.add_paragraph("Managed a cross-functional team of 12 engineers", style="List Number")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(tmp_path, docx_bytes)
    outcome = extract_docx_text(artifact, settings)

    assert len(outcome.blocks) == 4
    assert outcome.blocks[0].block_type == BlockType.HEADING
    assert outcome.blocks[0].raw_text == "Senior Architect Resume"

    assert outcome.blocks[1].block_type == BlockType.PARAGRAPH
    assert outcome.blocks[1].raw_text == "Summary of qualifications and experience."

    assert outcome.blocks[2].block_type == BlockType.LIST_ITEM
    assert outcome.blocks[2].raw_text == "Built cloud platform with 99.99% reliability"

    assert outcome.blocks[3].block_type == BlockType.LIST_ITEM
    assert outcome.blocks[3].raw_text == "Managed a cross-functional team of 12 engineers"


# =========================================================================
# 3. Table cells receive row-major structural locations and deduplicate merged cells
# =========================================================================

def test_table_cells_and_merged_cells(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    # Merge row 0 cell 0 and cell 1
    cell_a = table.cell(0, 0)
    cell_b = table.cell(0, 1)
    cell_a.merge(cell_b)
    cell_a.text = "Merged Header"

    table.cell(1, 0).text = "Row 1 Left"
    table.cell(1, 1).text = "Row 1 Right"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(tmp_path, docx_bytes)
    outcome = extract_docx_text(artifact, settings)

    # Merged cell is only extracted once
    assert len(outcome.blocks) == 3
    assert outcome.blocks[0].raw_text == "Merged Header"
    assert outcome.blocks[0].structural_location == "body/table/0/row/0/cell/0"

    assert outcome.blocks[1].raw_text == "Row 1 Left"
    assert outcome.blocks[1].structural_location == "body/table/0/row/1/cell/0"

    assert outcome.blocks[2].raw_text == "Row 1 Right"
    assert outcome.blocks[2].structural_location == "body/table/0/row/1/cell/1"


# =========================================================================
# 4. Visible hyperlink text is retained without opening target URL
# =========================================================================

def test_visible_hyperlink_text_retained(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Check my ")
    _add_hyperlink(p, "https://portfolio.example.com/candidate", "Portfolio Website")
    p.add_run(" for projects.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(tmp_path, docx_bytes)
    outcome = extract_docx_text(artifact, settings)

    assert len(outcome.blocks) == 1
    assert outcome.blocks[0].raw_text == "Check my Portfolio Website for projects."
    assert "https://" not in outcome.blocks[0].raw_text


# =========================================================================
# 5. Headers and footers are extracted without linked-part duplication
# =========================================================================

def test_headers_and_footers_deduplicated_across_sections(tmp_path):
    doc = Document()
    section1 = doc.sections[0]
    section1.header.paragraphs[0].text = "Header Text"
    section1.footer.paragraphs[0].text = "Footer Text"
    doc.add_paragraph("Body Page 1")

    # Add section 2 with linked header and footer
    section2 = doc.add_section()
    doc.add_paragraph("Body Page 2")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(tmp_path, docx_bytes)
    outcome = extract_docx_text(artifact, settings)

    # Expected order:
    # 0: header
    # 1: Body Page 1
    # 2: Body Page 2
    # 3: footer
    assert len(outcome.blocks) == 4
    assert outcome.blocks[0].block_type == BlockType.HEADER
    assert outcome.blocks[0].raw_text == "Header Text"
    assert outcome.blocks[0].structural_location == "section/0/header/default/paragraph/0"

    assert outcome.blocks[1].block_type == BlockType.PARAGRAPH
    assert outcome.blocks[1].raw_text == "Body Page 1"

    assert outcome.blocks[2].block_type == BlockType.PARAGRAPH
    assert outcome.blocks[2].raw_text == "Body Page 2"

    assert outcome.blocks[3].block_type == BlockType.FOOTER
    assert outcome.blocks[3].raw_text == "Footer Text"
    assert outcome.blocks[3].structural_location == "section/0/footer/default/paragraph/0"


# =========================================================================
# 6. DOCX blocks have null page numbers and valid structural locations
# =========================================================================

def test_docx_blocks_have_null_page_number_and_valid_structural_locations(tmp_path):
    doc = Document()
    doc.add_paragraph("Sample Paragraph")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell Content"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(tmp_path, docx_bytes)
    outcome = extract_docx_text(artifact, settings)

    for block in outcome.blocks:
        assert block.page_number is None
        assert block.structural_location is not None
        assert block.structural_location.startswith("body/")
        assert block.extraction_method == ExtractionMethod.DOCX_TEXT
        assert block.confidence is None


# =========================================================================
# 7. Empty/image-only DOCX returns a controlled unavailable outcome
# =========================================================================

def test_empty_docx_raises_text_unavailable(tmp_path):
    doc = Document()
    # Empty document without text
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(tmp_path, docx_bytes)

    with pytest.raises(DOCXTextUnavailableError) as exc:
        extract_docx_text(artifact, settings)
    assert "No native text" in str(exc.value)


# =========================================================================
# 8. Hash mismatch, unsafe containment, and extracted-text limits fail safely
# =========================================================================

def test_unvalidated_artifact_rejected(tmp_path):
    doc = Document()
    doc.add_paragraph("Valid content")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(
        tmp_path,
        docx_bytes,
        validation_status=ArtifactValidationStatus.PENDING,
    )

    with pytest.raises(ArtifactIntegrityError) as exc:
        extract_docx_text(artifact, settings)
    assert "VALIDATED state" in str(exc.value)


def test_hash_mismatch_rejected(tmp_path):
    doc = Document()
    doc.add_paragraph("Valid content")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(
        tmp_path,
        docx_bytes,
        hash_override="0" * 64,
    )

    with pytest.raises(ArtifactIntegrityError) as exc:
        extract_docx_text(artifact, settings)
    assert "hash mismatch" in str(exc.value)


def test_size_mismatch_rejected(tmp_path):
    doc = Document()
    doc.add_paragraph("Valid content")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(
        tmp_path,
        docx_bytes,
        size_override=999999,
    )

    with pytest.raises(ArtifactIntegrityError) as exc:
        extract_docx_text(artifact, settings)
    assert "size mismatch" in str(exc.value)


def test_quarantine_containment_path_traversal_rejected(tmp_path):
    doc = Document()
    doc.add_paragraph("Valid content")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(
        tmp_path,
        docx_bytes,
        quarantine_key_override="docx/../../../etc/passwd",
    )

    with pytest.raises(ArtifactIntegrityError) as exc:
        extract_docx_text(artifact, settings)
    assert "Quarantine" in str(exc.value)


def test_symlink_quarantined_file_rejected(tmp_path):
    doc = Document()
    doc.add_paragraph("Valid content")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(tmp_path, docx_bytes)

    # Replace real file with a symlink
    quarantine_base = settings.get_quarantine_path()
    real_file = quarantine_base / artifact.quarantine_key
    target_file = tmp_path / "other_target.docx"
    target_file.write_bytes(docx_bytes)

    real_file.unlink()
    real_file.symlink_to(target_file)

    with pytest.raises(ArtifactIntegrityError) as exc:
        extract_docx_text(artifact, settings)
    assert "symlink" in str(exc.value)


def test_extracted_character_limit_exceeded(tmp_path):
    doc = Document()
    doc.add_paragraph("A" * 500)
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(tmp_path, docx_bytes)
    settings.max_docx_extracted_chars = 100

    with pytest.raises(DOCXTextLimitExceededError) as exc:
        extract_docx_text(artifact, settings)
    assert "character limit" in str(exc.value)


# =========================================================================
# 9. Errors do not leak sensitive text or XML details
# =========================================================================

def test_corrupted_docx_file_returns_controlled_error_without_leaks(tmp_path):
    corrupt_bytes = b"PK\x03\x04Corrupted DOCX with SECRET_TOKEN_CONFIDENTIAL_12345"
    artifact, settings = _setup_quarantined_docx(tmp_path, corrupt_bytes)

    with pytest.raises(DOCXExtractionFailedError) as exc:
        extract_docx_text(artifact, settings)

    err_msg = str(exc.value)
    assert "SECRET_TOKEN" not in err_msg
    assert "CONFIDENTIAL" not in err_msg
    assert "Failed to" in err_msg


# =========================================================================
# 10. A valid result satisfies and persists through Phase 1C.1 contract
# =========================================================================

class MockDBSession:
    def __init__(self):
        self.added = []
        self.flushed = False

    def add(self, instance):
        self.added.append(instance)

    def flush(self):
        self.flushed = True

    def execute(self, stmt):
        mock_result = MagicMock()
        mock_result.scalar_one_or_none.return_value = None
        return mock_result


def test_valid_docx_result_persists_with_phase_1c1_service(tmp_path):
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("Principal Cloud Architect")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    artifact, settings = _setup_quarantined_docx(tmp_path, docx_bytes)
    outcome = extract_docx_text(artifact, settings)
    assert outcome.extraction_result is not None

    db_session = MockDBSession()
    persisted = persist_extraction(
        session=db_session,
        artifact=artifact,
        result=outcome.extraction_result,
    )

    assert isinstance(persisted, ResumeExtraction)
    assert persisted.extractor_version == "docx-structural-text-v1.0"
    assert len(persisted.blocks) == 2
    assert persisted.blocks[0].sequence_index == 0
    assert persisted.blocks[0].page_number is None
    assert persisted.blocks[0].structural_location == "body/paragraph/0"
    assert persisted.blocks[0].raw_text == "Jane Doe"
    assert persisted.blocks[0].extraction_method == ExtractionMethod.DOCX_TEXT
    assert persisted.blocks[1].sequence_index == 1
    assert persisted.blocks[1].structural_location == "body/paragraph/1"
    assert persisted.blocks[1].raw_text == "Principal Cloud Architect"
    assert db_session.flushed is True
