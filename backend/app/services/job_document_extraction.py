import io
import logging
import os
import tempfile
from pathlib import Path
from typing import Tuple

import docx
import pypdf

from app.models.resume import ArtifactFormat
from app.services.docx_validation import DOCXValidationError, validate_docx_file
from app.services.pdf_validation import PDFValidationError, validate_pdf_file

logger = logging.getLogger("ai_hiring.services.job_document_extraction")

MAX_JD_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_JD_PDF_PAGES = 50


class JobDocumentError(Exception):
    def __init__(self, code: str, message: str, status_code: int = 422):
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code


def detect_document_format(file_bytes: bytes, filename: str) -> ArtifactFormat:
    """Detects document format from magic bytes and filename."""
    if len(file_bytes) < 4:
        raise JobDocumentError("EMPTY_OR_CORRUPT_FILE", "The uploaded file is empty or corrupted.")

    if file_bytes.startswith(b"%PDF"):
        return ArtifactFormat.pdf

    if file_bytes.startswith(b"PK\x03\x04"):
        # ZIP archive -> check if DOCX
        if filename.lower().endswith(".docx"):
            return ArtifactFormat.docx

    raise JobDocumentError(
        "UNSUPPORTED_DOCUMENT_FORMAT",
        "Unsupported document format. Only secure PDF (.pdf) and Word (.docx) documents are supported. Alternatively, you can paste the text directly.",
    )


def extract_text_from_document(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """
    Validates and extracts raw text from an uploaded PDF or DOCX file.
    Returns (extracted_text, detected_format_str).
    """
    if len(file_bytes) > MAX_JD_FILE_SIZE:
        raise JobDocumentError(
            "FILE_SIZE_EXCEEDED",
            f"File exceeds maximum allowed size of {MAX_JD_FILE_SIZE // (1024 * 1024)} MB.",
        )

    detected_format = detect_document_format(file_bytes, filename)

    # Write temporarily to run security validation & direct text extraction
    with tempfile.NamedTemporaryFile(suffix=f".{detected_format.value.lower()}", delete=False) as tmp:
        tmp_path = Path(tmp.name)
        tmp.write(file_bytes)

    try:
        if detected_format == ArtifactFormat.pdf:
            # 1. Validate PDF security (checks actions, js, annotations, launch, etc.)
            try:
                validate_pdf_file(tmp_path, max_pages=MAX_JD_PDF_PAGES)
            except PDFValidationError as e:
                raise JobDocumentError(e.code, f"PDF security validation failed: {e.message}")

            # 2. Extract text with pypdf
            try:
                reader = pypdf.PdfReader(str(tmp_path))
                extracted_pages = []
                for page in reader.pages:
                    txt = page.extract_text() or ""
                    if txt.strip():
                        extracted_pages.append(txt.strip())
            except Exception as e:
                raise JobDocumentError("PDF_EXTRACTION_ERROR", f"Failed to extract text from PDF: {str(e)}")

            if not extracted_pages:
                raise JobDocumentError(
                    "IMAGE_ONLY_DOCUMENT_NO_OCR",
                    "The uploaded PDF contains no extractable text (image-only or scanned document). OCR is not supported. Please paste the job description text manually.",
                )

            extracted_text = "\n\n".join(extracted_pages)
            return extracted_text, "PDF"

        elif detected_format == ArtifactFormat.docx:
            # 1. Validate DOCX security (checks xml bombs, macros, corrupted zips, etc.)
            try:
                validate_docx_file(tmp_path)
            except DOCXValidationError as e:
                raise JobDocumentError(e.code, f"DOCX security validation failed: {e.message}")

            # 2. Extract text from DOCX
            try:
                doc = docx.Document(str(tmp_path))
                paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            c_txt = cell.text.strip()
                            if c_txt and c_txt not in paragraphs:
                                paragraphs.append(c_txt)
            except Exception as e:
                raise JobDocumentError("DOCX_EXTRACTION_ERROR", f"Failed to extract text from DOCX: {str(e)}")

            if not paragraphs:
                raise JobDocumentError(
                    "IMAGE_ONLY_DOCUMENT_NO_OCR",
                    "The uploaded DOCX contains no extractable text. Please paste the job description text manually.",
                )

            extracted_text = "\n\n".join(paragraphs)
            return extracted_text, "DOCX"

        else:
            raise JobDocumentError("UNSUPPORTED_FORMAT", "Unsupported document format.")

    finally:
        if tmp_path.exists():
            try:
                os.unlink(tmp_path)
            except Exception:
                pass
